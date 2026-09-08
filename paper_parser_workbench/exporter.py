"""
文档导出模块
功能：导出规范的docx文献阅读笔记
"""
import io
from datetime import datetime
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from typing import Dict, List, Tuple


def create_document():
    """创建新的Word文档"""
    doc = Document()
    
    # 设置默认字体
    style = doc.styles['Normal']
    font = style.font
    font.name = '宋体'
    font.size = Pt(12)
    
    return doc


def add_title(doc: Document, title: str):
    """添加文档标题"""
    # 添加标题
    heading = doc.add_heading(level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    run = heading.add_run(title)
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(0xC8, 0x10, 0x2E)  # 师大红


def add_metadata_section(doc: Document, metadata: Dict[str, str]):
    """添加元数据部分"""
    # 添加分隔线
    doc.add_paragraph('─' * 50)
    
    # 标题
    heading = doc.add_heading('文档信息', level=1)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0xC8, 0x10, 0x2E)
    
    # 元数据内容
    fields = [
        ("文献标题", metadata.get('title', '未识别')),
        ("作者", metadata.get('authors', '未识别')),
        ("关键词", metadata.get('keywords', '未识别')),
    ]
    
    for label, value in fields:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        
        # 标签（加粗）
        label_run = p.add_run(f"【{label}】")
        label_run.bold = True
        label_run.font.size = Pt(12)
        
        # 内容
        content_run = p.add_run(f" {value}")
        content_run.font.size = Pt(12)


def add_abstract_section(doc: Document, abstract: str):
    """添加摘要部分"""
    if not abstract:
        return
    
    # 添加分隔线
    doc.add_paragraph('─' * 50)
    
    # 标题
    heading = doc.add_heading('摘要', level=1)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0xC8, 0x10, 0x2E)
    
    # 摘要内容
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.line_spacing = 1.5
    
    run = p.add_run(abstract)
    run.font.size = Pt(12)


def add_keywords_section(doc: Document, keywords: List[Tuple[str, float]]):
    """添加关键词部分"""
    if not keywords:
        return
    
    # 添加分隔线
    doc.add_paragraph('─' * 50)
    
    # 标题
    heading = doc.add_heading('高频关键词', level=1)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0xC8, 0x10, 0x2E)
    
    # 关键词列表
    for i, (word, weight) in enumerate(keywords[:15], 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        
        # 排名
        rank_run = p.add_run(f"{i}. ")
        rank_run.bold = True
        rank_run.font.size = Pt(12)
        
        # 关键词
        word_run = p.add_run(word)
        word_run.font.size = Pt(12)
        
        # 权重
        weight_run = p.add_run(f" (权重: {weight:.4f})")
        weight_run.font.size = Pt(10)
        weight_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def add_conclusions_section(doc: Document, conclusions: str):
    """添加核心结论部分"""
    if not conclusions:
        return
    
    # 添加分隔线
    doc.add_paragraph('─' * 50)
    
    # 标题
    heading = doc.add_heading('核心结论', level=1)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0xC8, 0x10, 0x2E)
    
    # 结论内容
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.line_spacing = 1.5
    
    run = p.add_run(conclusions)
    run.font.size = Pt(12)


def add_text_preview_section(doc: Document, text: str, max_length: int = 2000):
    """添加文本预览部分"""
    if not text:
        return
    
    # 添加分隔线
    doc.add_paragraph('─' * 50)
    
    # 标题
    heading = doc.add_heading('正文预览', level=1)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0xC8, 0x10, 0x2E)
    
    # 截取文本
    preview_text = text[:max_length]
    if len(text) > max_length:
        preview_text += "..."
    
    # 文本内容
    paragraphs = preview_text.split('\n')
    for para_text in paragraphs:
        if para_text.strip():
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Pt(24)
            p.paragraph_format.line_spacing = 1.5
            
            run = p.add_run(para_text)
            run.font.size = Pt(12)


def add_footer(doc: Document):
    """添加页脚信息"""
    # 添加分隔线
    doc.add_paragraph('─' * 50)
    
    # 生成信息
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    run = p.add_run(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    run2 = p2.add_run("学术文献解析与报告生成系统")
    run2.font.size = Pt(10)
    run2.font.color.rgb = RGBColor(0x99, 0x99, 0x99)


def export_to_docx(metadata: Dict[str, str],
                  abstract: str,
                  keywords: List[Tuple[str, float]],
                  conclusions: str,
                  text: str) -> io.BytesIO:
    """
    导出文献阅读笔记为docx格式
    参数：
        metadata: 文档元数据
        abstract: 摘要内容
        keywords: 关键词列表
        conclusions: 核心结论
        text: 正文文本
    返回：
        docx文件的BytesIO对象
    """
    # 创建文档
    doc = create_document()
    
    # 添加各个部分
    title = metadata.get('title', '文献阅读笔记')
    add_title(doc, title)
    add_metadata_section(doc, metadata)
    add_abstract_section(doc, abstract)
    add_keywords_section(doc, keywords)
    add_conclusions_section(doc, conclusions)
    add_text_preview_section(doc, text)
    add_footer(doc)
    
    # 保存到内存
    doc_buffer = io.BytesIO()
    doc.save(doc_buffer)
    doc_buffer.seek(0)
    
    return doc_buffer


def get_download_filename(metadata: Dict[str, str]) -> str:
    """生成下载文件名"""
    title = metadata.get('title', '文献阅读笔记')
    # 清理文件名中的非法字符
    safe_title = "".join(c for c in title if c.isalnum() or c in (' ', '-', '_'))
    safe_title = safe_title.strip()[:50]  # 限制长度
    
    if not safe_title:
        safe_title = "文献阅读笔记"
    
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    return f"{safe_title}_{timestamp}.docx"
